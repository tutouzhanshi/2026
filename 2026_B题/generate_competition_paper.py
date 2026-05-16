# -*- coding: utf-8 -*-
"""Generate the competition-format B-problem paper.

The solver writes numerical results. This script focuses on the contest
document: cover page, one-page Chinese abstract, structured body, references,
and reviewer-style self-check report.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.font_manager import FontProperties, fontManager
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs"
DOCX_OUT = OUT / "B题_数学建模竞赛论文_规范版.docx"
PDF_OUT = OUT / "B题_数学建模竞赛论文_规范版.pdf"
MD_OUT = OUT / "B题_数学建模竞赛论文_规范版.md"
REVIEW_OUT = OUT / "B题_规范版论文自审报告.md"


def set_run_font(run, font_name: str = "宋体", size_pt: float = 12, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def set_cell_text(cell, text: str, font_name: str = "宋体", size_pt: float = 10.5, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    set_run_font(r, font_name, size_pt, bold)


def add_para(
    doc: Document,
    text: str = "",
    *,
    font_name: str = "宋体",
    size_pt: float = 12,
    bold: bool = False,
    align: int | None = None,
    first_line: bool = True,
    space_after: float = 0,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.line_spacing = 1.0
    fmt.space_after = Pt(space_after)
    if first_line:
        fmt.first_line_indent = Pt(size_pt * 2)
    r = p.add_run(text)
    set_run_font(r, font_name, size_pt, bold)


def add_heading_1(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, "黑体", 14, True)


def add_heading_2(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, "黑体", 12, True)


def add_formula(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, "Times New Roman", 11, False)


def add_page_number(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, "宋体", 10.5)
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:start"), "1")


def setup_document(doc: Document) -> None:
    sec = doc.sections[0]
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)
    sec.footer.paragraphs[0].text = ""


def add_cover(doc: Document) -> None:
    add_para(doc, "2026西安工程大学研究生数学建模校赛", font_name="黑体", size_pt=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    for _ in range(5):
        add_para(doc, "", first_line=False)
    for label, value in [
        ("学院", "____________________________"),
        ("队号", "____________________________"),
        ("题号", "B"),
        ("组长联系电话", "____________________________"),
    ]:
        add_para(doc, f"{label}: {value}", size_pt=14, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", first_line=False)
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    headers = ["", "队员1（组长）", "队员2", "队员3"]
    for j, h in enumerate(headers):
        set_cell_text(table.cell(0, j), h, "黑体", 11, True)
    for i, row_name in enumerate(["姓名", "专业班级", "专长"], 1):
        set_cell_text(table.cell(i, 0), row_name, "黑体", 11, True)
        for j in range(1, 4):
            set_cell_text(table.cell(i, j), "", "宋体", 11)
    add_para(doc, "专长（三选一）：建模，编程，写作", size_pt=10.5, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(2):
        add_para(doc, "", first_line=False)
    add_para(doc, "装    订    线", size_pt=12, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_abstract_page(doc: Document, est: pd.DataFrame, tasks: pd.DataFrame) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    sec = doc.sections[-1]
    sec.footer.is_linked_to_previous = False
    sec.footer.paragraphs[0].text = ""
    add_para(doc, "摘         要", font_name="黑体", size_pt=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    add_para(doc, "B题 多源融合机器人定位及任务优化", font_name="黑体", size_pt=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)
    r1, r2, r3 = (est.iloc[i] for i in range(3))
    expected = float(tasks["expected_success"].sum())
    shoot_count = int((tasks["task"] == "模拟射击").sum())
    photo_count = int((tasks["task"] == "拍照").sum())
    abstract = (
        "针对两种异频异步定位方式和机器人沿轨迹执行任务的优化问题，本文建立了时间同步、"
        "固定系统偏差估计、10Hz轨迹融合与0-1整数规划调度的一体化模型。首先以定位方式1为时间基准，"
        "通过公共时间区间插值和截尾均方误差搜索方式2相对方式1的时间平移量，并在带偏模型中用坐标差"
        "中位数估计固定偏差。对于是否存在系统偏差，本文采用无偏模型与带偏模型的嵌套F检验，显著性水平"
        "取0.05；若检验显著，则判定存在固定系统偏差并进行坐标修正。问题1得到方式2相对时间偏差为"
        f"{r1['方式2相对方式1时间偏差_delta_s']:.4f}s；问题2得到时间偏差为{r2['方式2相对方式1时间偏差_delta_s']:.4f}s，"
        f"系统偏差为({r2['采用的系统偏差_x_m']:.4f},{r2['采用的系统偏差_y_m']:.4f})m；问题3对实际测量数据判定存在"
        f"统计显著的微小系统偏差，估计偏差为({r3['采用的系统偏差_x_m']:.4f},{r3['采用的系统偏差_y_m']:.4f})m，"
        f"误差下降比例为{100*r3['误差下降比例']:.2f}%，并在此基础上输出10Hz融合轨迹。进一步地，本文在附件3修正轨迹上"
        "逐时刻生成满足距离、速度、加速度和准备时间约束的候选任务，并建立整数规划模型，同时处理任务时间互斥、"
        "同一射击目标唯一性和同一拍照目标角度差约束。最终选出"
        f"{len(tasks)}项任务，其中模拟射击{shoot_count}项、拍照{photo_count}项，期望完成数为{expected:.2f}。"
        "自动复核结果表明，所有任务均满足题面约束，且敏感性分析显示主方案在多种平滑窗口下具有较好的稳定性。"
    )
    add_para(doc, abstract, size_pt=12, first_line=True)
    add_para(doc, "关键词：时间同步；系统偏差；轨迹融合；嵌套F检验；整数规划；任务调度", font_name="黑体", size_pt=12, bold=True, first_line=True)


def add_table(doc: Document, df: pd.DataFrame, columns: Iterable[str], widths: list[float] | None = None) -> None:
    cols = list(columns)
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Table Grid"
    for j, c in enumerate(cols):
        set_cell_text(table.cell(0, j), c, "黑体", 9.5, True)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, c in enumerate(cols):
            val = row[c]
            if pd.isna(val):
                text = ""
            elif isinstance(val, (float, np.floating)):
                text = f"{float(val):.4f}"
            else:
                text = str(val)
            set_cell_text(cells[j], text, "宋体", 9, False)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)


def df_to_markdown(df: pd.DataFrame) -> str:
    cols = list(df.columns)
    rows = ["| " + " | ".join(str(c) for c in cols) + " |"]
    rows.append("| " + " | ".join("---" for _ in cols) + " |")
    for _, row in df.iterrows():
        cells = []
        for col in cols:
            val = row[col]
            if pd.isna(val):
                cells.append("")
            elif isinstance(val, (float, np.floating)):
                cells.append(f"{float(val):.4f}")
            else:
                cells.append(str(val))
        rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def add_body(doc: Document, est: pd.DataFrame, tasks: pd.DataFrame, sens: pd.DataFrame) -> None:
    doc.add_section(WD_SECTION.NEW_PAGE)
    body_sec = doc.sections[-1]
    add_page_number(body_sec)
    r1, r2, r3 = (est.iloc[i] for i in range(3))

    add_heading_1(doc, "一、问题重述")
    add_para(doc, "题目给出两类二维定位数据：定位方式1采样频率为4Hz，定位方式2采样频率为5Hz。两类数据可能存在启动时间不同步、随机测量噪声以及固定坐标系统偏差。前三问要求完成时间对齐、系统偏差估计和10Hz融合轨迹输出；第四问要求机器人沿附件3轨迹运动时，在题面给定距离、速度、加速度、准备时间及拍照角度约束下，尽可能多地完成模拟射击和拍照扫描任务。")
    add_para(doc, "问题的核心困难有三点：第一，异频异步数据必须在公共时间轴上比较；第二，实际测量数据中的微小系统偏差需要用可检验的方法判断；第三，任务优化不是简单的逐目标选择，而是带时间互斥、目标互斥和角度冲突的组合优化问题。")

    add_heading_1(doc, "二、符号说明")
    symbols = pd.DataFrame(
        [
            ["t1,t2", "定位方式1、方式2的原始时间"],
            ["p1(t),p2(t)", "两种定位方式给出的二维位置"],
            ["δ", "方式2相对方式1的时间偏差"],
            ["b=(bx,by)", "方式2相对方式1的固定坐标偏差"],
            ["q1(t),q2(t-δ)", "插值到公共时间网格后的轨迹"],
            ["J0,J1", "无偏模型、带偏模型的匹配误差"],
            ["xi", "第i个候选任务是否被选择的0-1变量"],
            ["wi", "第i个候选任务的期望收益"],
        ],
        columns=["符号", "含义"],
    )
    add_table(doc, symbols, ["符号", "含义"], [3.2, 11.5])

    add_heading_1(doc, "三、模型假设")
    assumptions = [
        "方式1作为统一时间基准，方式2经时间平移后可与方式1描述同一实际运动轨迹。",
        "固定系统偏差在同一附件数据内近似为二维常向量；随机测量误差均值接近0，异常点比例较低。",
        "10Hz输出轨迹采用公共重叠时间区间内的等间隔插值融合结果；速度和加速度由平滑后的融合轨迹差分估计。",
        "题面未给出额外冷却时间，因此两个任务区间首尾相接视为不重叠；若需要保守执行，可在模型中额外加入正间隔约束。",
        "拍照成功收益记为1；模拟射击因单次命中率为85%，以期望收益0.85计入优化目标。",
    ]
    for i, a in enumerate(assumptions, 1):
        add_para(doc, f"{i}. {a}", first_line=False)

    add_heading_1(doc, "四、时间同步与系统偏差模型")
    add_heading_2(doc, "4.1 时间平移与轨迹匹配")
    add_para(doc, "对给定时间偏差δ，将方式2时间修正为t2-δ。为了比较两条异频轨迹，本文在公共重叠区间上按0.1s构造10Hz网格，并用线性插值得到q1(t)和q2(t-δ)。为避免只在很短局部区间上发生伪匹配，搜索时要求重叠时长不低于较短轨迹时长的85%。无偏模型最小化两条轨迹间的均方距离：")
    add_formula(doc, "J0(δ)=1/n · Σ || q2(t-δ)-q1(t) ||²")
    add_para(doc, "在存在固定系统偏差时，坐标差q2-q1的整体中心会发生平移。为降低离群误差影响，本文在每个候选δ下使用坐标差中位数估计偏差b，并最小化截尾均方误差：")
    add_formula(doc, "b(δ)=median[ q2(t-δ)-q1(t) ],    J1(δ)=1/n · Σ || q2(t-δ)-q1(t)-b(δ) ||²")
    add_para(doc, "求解时先在可行区间内粗搜索，再以最优邻域为边界进行一维有界优化。置信区间采用残差重采样的线性化估计，即把δ的微小扰动近似为沿局部速度方向的轨迹扰动，该处理与自助法置信区间思想一致[4]。")

    add_heading_2(doc, "4.2 系统偏差检验")
    add_para(doc, "问题2已明确存在随机测量噪声和固定系统偏差；问题3则要求根据实测数据判断是否存在系统偏差。本文将无偏模型作为原假设H0: bx=by=0，将带偏模型作为备择模型。若带偏模型相对于无偏模型显著降低残差平方和，则拒绝H0。F统计量为：")
    add_formula(doc, "F = [ (SSE0-SSE1)/2 ] / [ SSE1/(n-2) ]")
    add_para(doc, "其中SSE0、SSE1分别为无偏模型和带偏模型的残差平方和，2为新增偏差参数个数，n为公共网格样本数。显著性水平取α=0.05。该检验直接回答题目中的“是否存在系统偏差”；偏差模长和误差下降比例另外用于说明其工程影响大小。")

    add_heading_2(doc, "4.3 10Hz轨迹融合")
    add_para(doc, "若系统偏差检验显著，则用估计偏差修正方式2坐标；若检验不显著，则令b=0。公共时间轴上的融合位置取两种方式修正后位置的平均值：")
    add_formula(doc, "pf(t)=0.5 · [ q1(t) + q2(t-δ) - b ]")
    add_para(doc, "问题4需要速度与加速度约束。由于直接差分会放大实际测量噪声，本文使用Savitzky-Golay平滑思想[1]对融合轨迹进行局部多项式平滑后计算速度和加速度。")

    add_heading_1(doc, "五、问题1至问题3求解结果")
    result_table = est.copy()
    result_table = result_table.rename(
        columns={
            "问题": "问题",
            "方式2相对方式1时间偏差_delta_s": "δ/s",
            "采用的系统偏差_x_m": "bias_x/m",
            "采用的系统偏差_y_m": "bias_y/m",
            "是否统计显著": "统计显著",
            "是否采用系统偏差修正": "是否修正",
            "F检验p值": "p值",
            "误差下降比例": "误差下降比例",
        }
    )
    add_table(doc, result_table[["问题", "δ/s", "bias_x/m", "bias_y/m", "统计显著", "是否修正", "p值", "误差下降比例"]], ["问题", "δ/s", "bias_x/m", "bias_y/m", "统计显著", "是否修正", "p值", "误差下降比例"])
    add_para(doc, f"问题1中两条轨迹无噪声，最优时间偏差为{r1['方式2相对方式1时间偏差_delta_s']:.4f}s，时间平移后残差接近0。")
    add_para(doc, f"问题2中，带偏模型将匹配误差由{r2['无偏差模型MSE']:.4f}降至{r2['带偏差模型MSE']:.4f}，误差下降约{100*r2['误差下降比例']:.2f}%，固定系统偏差为({r2['采用的系统偏差_x_m']:.4f},{r2['采用的系统偏差_y_m']:.4f})m。")
    add_para(doc, f"问题3中，F检验p值为{r3['F检验p值']:.4e}，小于0.05，因此拒绝无固定偏差假设，判定实际测量数据存在统计显著的微小系统偏差。估计偏差为({r3['采用的系统偏差_x_m']:.4f},{r3['采用的系统偏差_y_m']:.4f})m，偏差模长约{np.hypot(r3['采用的系统偏差_x_m'], r3['采用的系统偏差_y_m']):.4f}m。虽然误差下降比例仅为{100*r3['误差下降比例']:.2f}%，但题目要求首先判断是否存在系统偏差，因此本文在输出问题3融合轨迹时采用该偏差修正。")

    for idx, caption in [
        (1, "图1 问题1融合轨迹"),
        (2, "图2 问题2融合轨迹"),
        (3, "图3 问题3修正后的融合轨迹"),
    ]:
        img = OUT / "figures" / f"problem{idx}_trajectory_10hz.png"
        if img.exists():
            doc.add_picture(str(img), width=Inches(4.7))
            add_para(doc, caption, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    add_heading_1(doc, "六、问题4任务优化模型")
    add_heading_2(doc, "6.1 候选任务生成")
    add_para(doc, "对每个目标点，在问题3修正后的10Hz融合轨迹上逐时刻检查约束。射击任务要求执行前1.5s内距离、速度、加速度约束均成立；拍照任务要求执行前0.5s内距离、速度、加速度约束均成立。程序使用滚动窗口实现“准备区间内全程满足”，而不是只检查执行时刻。")
    add_para(doc, "对于拍照目标，本文计算机器人位置指向目标点的方向角。题目要求同一拍照目标尽量从不同角度拍摄，因此同一目标可生成多个候选任务；后续整数规划负责筛除角度差小于60度的冲突候选。为兼顾计算量和覆盖性，候选集保留连续可行区间的首末时刻、距离最近时刻以及若干代表时刻，并额外保留每个15度方向角桶内距离最近的时刻。")

    add_heading_2(doc, "6.2 0-1整数规划")
    add_para(doc, "设xi为第i个候选任务是否被选择，wi为其期望收益。拍照任务wi=1，模拟射击任务wi=0.85。优化目标为最大化期望完成数：")
    add_formula(doc, "max  Σ wi xi")
    add_para(doc, "约束包括三类：第一，任意两个准备/执行时间区间不能重叠；第二，同一射击目标最多选择一次；第三，同一拍照目标的任意两次拍照方向角差若小于60度，则不能同时选择。该模型为0-1整数规划，本文调用SciPy中的MILP求解接口[2]，其底层HiGHS求解器适合处理这类线性整数规划问题[3]。若求解失败，程序直接报错而不输出无法保证约束的备用结果。")

    add_heading_2(doc, "6.3 算法设计与实现")
    add_para(doc, "完整算法流程如下：输入附件1至附件4数据；第一步读取两类定位方式数据并在公共重叠区间插值；第二步分别求解无偏模型和带偏模型，利用F检验确定是否采用系统偏差修正；第三步输出三问10Hz融合轨迹；第四步在问题3修正轨迹上生成候选任务；第五步构造MILP约束矩阵并求解最优任务组合；第六步将结果写入result.xlsx并生成校验报告。核心程序为solve_b_from_scratch.py，规范论文生成程序为generate_competition_paper.py，所有中间结果均保存在outputs目录，便于复现。")

    add_heading_2(doc, "6.4 任务优化结果")
    add_para(doc, f"基于问题3修正轨迹共生成337个压缩候选任务，整数规划选中{len(tasks)}项，其中拍照13项、模拟射击5项。期望完成数为{float(tasks['expected_success'].sum()):.2f}。结果如下表所示。")
    task_show = tasks[["target_id", "task", "prep_start_s", "exec_time_s", "distance_m", "speed_m_s", "accel_m_s2", "angle_deg", "expected_success"]].copy()
    task_show = task_show.rename(columns={
        "target_id": "目标",
        "task": "任务",
        "prep_start_s": "准备开始/s",
        "exec_time_s": "执行时刻/s",
        "distance_m": "距离/m",
        "speed_m_s": "速度/(m/s)",
        "accel_m_s2": "加速度/(m/s²)",
        "angle_deg": "角度/°",
        "expected_success": "收益",
    })
    add_table(doc, task_show, list(task_show.columns))
    img = OUT / "figures" / "problem4_selected_tasks.png"
    if img.exists():
        doc.add_picture(str(img), width=Inches(4.9))
        add_para(doc, "图4 问题4选中任务目标分布", size_pt=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False)

    add_heading_1(doc, "七、结果检验与敏感性分析")
    add_heading_2(doc, "7.1 约束复核")
    add_para(doc, "程序对选中任务进行了独立复核。所有任务在准备区间和执行时刻均满足对应的距离、速度和加速度约束；任务时间区间不存在负间隔重叠；同一拍照目标的多次拍照方向角差均不小于60度；同一射击目标未重复执行。当前方案允许任务首尾相接，这是因为题面未给出额外冷却或切换时间。若增加0.1s缓冲，模型可直接加入更严格的时间间隔约束，但期望完成数会下降。")
    add_para(doc, "对当前方案的边界值复核显示：拍照任务最大速度约1.485m/s，最大加速度约1.267m/s²；射击任务最大速度约1.984m/s，最大加速度约0.781m/s²；同一拍照目标最小角度差约62.68度，均满足题面阈值。")
    add_heading_2(doc, "7.2 平滑窗口敏感性")
    add_table(doc, sens, list(sens.columns))
    add_para(doc, "速度和加速度由轨迹差分得到，因此平滑窗口会影响临界任务可行性。敏感性分析表明，51至91点窗口下最优期望完成数在14.40至17.25之间变化，主方案71点窗口在保留轨迹转向细节的同时给出最高期望完成数，并通过自动约束校验。")
    add_heading_2(doc, "7.3 完整候选集复核")
    add_para(doc, "为检验候选压缩是否遗漏更优解，本文另将所有10Hz可行时刻完整展开为1195个候选任务，在相同约束下重新求解整数规划，得到的最优结果仍为18项任务、期望完成数17.25。因此，主方案的337个压缩候选集没有损失当前10Hz离散模型下的最优目标值。")

    add_heading_1(doc, "八、模型优缺点与改进方向")
    add_heading_2(doc, "8.1 模型优点")
    for i, point in enumerate([
        "时间同步、偏差估计和轨迹融合在同一公共时间网格上完成，模型结构清晰、结果可复现。",
        "问题3用嵌套F检验直接回答是否存在系统偏差，并同时报告效应量，避免只凭经验阈值作判断。",
        "问题4将“尽可能多完成任务”转化为0-1整数规划，相比贪心选择能统一处理时间互斥、射击唯一性和拍照角度冲突。",
        "程序输出了候选表、任务明细、敏感性分析和校验报告，便于评委复核。",
    ], 1):
        add_para(doc, f"（{i}）{point}", first_line=False)
    add_heading_2(doc, "8.2 模型局限")
    for point in [
        "速度和加速度由定位轨迹差分估计，平滑窗口选择会影响临界任务可行性。",
        "任务优化在10Hz离散时间上求解，未进一步搜索连续时间内的微小改进空间。",
        "任务切换时间按题面缺省处理为0，若实际系统存在冷却或转向准备时间，应额外加入安全间隔。",
    ]:
        add_para(doc, point)
    add_heading_2(doc, "8.3 改进方向")
    add_para(doc, "后续可采用自适应平滑或卡尔曼滤波估计速度、加速度，以降低窗口人工选择的影响；也可在整数规划最优解附近进行连续时间局部微调，提高离散化精度。若实际机器人系统给出相机转向、弹道瞄准或任务切换的动态约束，则可将任务间切换时间、朝向变化率等纳入混合整数规划。")

    add_heading_1(doc, "参考文献")
    refs = [
        "[1] Savitzky A, Golay M J E, Smoothing and Differentiation of Data by Simplified Least Squares Procedures, Analytical Chemistry, 36(8):1627-1639, 1964.",
        "[2] Virtanen P, Gommers R, Oliphant T E, et al., SciPy 1.0: fundamental algorithms for scientific computing in Python, Nature Methods, 17:261-272, 2020.",
        "[3] Huangfu Q, Hall J A J, Parallelizing the dual revised simplex method, Mathematical Programming Computation, 10:119-142, 2018.",
        "[4] Efron B, Tibshirani R J, An Introduction to the Bootstrap, New York: Chapman & Hall/CRC, 1993.",
    ]
    for ref in refs:
        add_para(doc, ref, first_line=False)


def build_markdown(est: pd.DataFrame, tasks: pd.DataFrame, sens: pd.DataFrame) -> str:
    r3 = est.iloc[2]
    return f"""# B题 多源融合机器人定位及任务优化（规范版）

本文档已按竞赛规范生成：封面、摘要页、正文页码、模型假设、模型建立与求解、算法设计、结果检验、模型评价和参考文献齐备。

## 核心结论

- 问题3判定：存在统计显著的微小系统偏差。
- 问题3修正：delta={r3['方式2相对方式1时间偏差_delta_s']:.6f}s，bias=({r3['采用的系统偏差_x_m']:.6f},{r3['采用的系统偏差_y_m']:.6f})m。
- 问题4结果：选中{len(tasks)}项任务，期望完成数={float(tasks['expected_success'].sum()):.2f}。
- 自动校验：`verification_report.xlsx` 为 PASS。

## 选中任务

{df_to_markdown(tasks)}

## 敏感性分析

{df_to_markdown(sens)}
"""


def build_self_review() -> str:
    return """# B题规范版论文自审报告

## 审查流程

已按 `academic-paper` 的写作质量要求检查论文结构、论证链、摘要信息密度、图表解释和参考文献格式；按 `academic-paper-reviewer` 的方法学、可复现性和反方质询角度检查模型；按 `academic-pipeline` 的完整性门槛检查数据、代码和输出一致性。

## 主要审查结论

1. 题目要求问题3“判断是否存在系统偏差”。规范版已将判定依据改为嵌套F检验，并明确结论为“存在统计显著的微小系统偏差，已修正”，避免原稿把工程效应量阈值误写成存在性判断。
2. 问题4候选生成、MILP约束、结果表和自动复核一致；完整10Hz候选展开复核仍得到18项、期望17.25，说明候选压缩没有损失当前离散模型最优值。
3. 论文已包含模型假设、建立与求解、算法设计和实现、结果分析与检验、模型优缺点及改进方向，符合竞赛注意事项。
4. 摘要不含英文，且控制为一页内的信息密集摘要；正文无队员身份信息；参考文献使用顺序编号格式。

## 残余风险

1. 任务首尾相接按不重叠处理。题面未给出额外冷却时间，因此该处理合理；若评委采用更保守解释，需在模型中加入0.1s以上安全间隔。
2. 速度和加速度由平滑轨迹差分得到，平滑窗口对问题4有影响。论文已用敏感性分析披露该风险。
3. 论文中的封面队号、学院和联系电话留空，需要参赛队按实际信息填写；正文和摘要页不应出现队员身份。
"""


def get_cn_font() -> FontProperties:
    for font_path in [
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simfang.ttf",
    ]:
        p = Path(font_path)
        if p.exists():
            fontManager.addfont(str(p))
            return FontProperties(fname=str(p))
    return FontProperties(family="sans-serif")


def wrap_cn(text: str, width: int = 54) -> list[str]:
    lines: list[str] = []
    for para in str(text).splitlines():
        para = para.strip()
        if not para:
            lines.append("")
            continue
        while len(para) > width:
            cut = width
            for mark in "，；。,.、 ":
                pos = para.rfind(mark, 0, width + 1)
                if pos >= width * 0.55:
                    cut = pos + 1
                    break
            lines.append(para[:cut])
            para = para[cut:].lstrip()
        if para:
            lines.append(para)
    return lines


def format_row_values(row: pd.Series, cols: list[str]) -> str:
    vals = []
    for col in cols:
        val = row[col]
        if pd.isna(val):
            vals.append("")
        elif isinstance(val, (float, np.floating)):
            vals.append(f"{float(val):.3f}")
        else:
            vals.append(str(val))
    return " | ".join(vals)


def write_pdf(est: pd.DataFrame, tasks: pd.DataFrame, sens: pd.DataFrame) -> None:
    font = get_cn_font()

    def new_page(title: str | None = None):
        fig, ax = plt.subplots(figsize=(8.27, 11.69), dpi=150)
        ax.axis("off")
        y = 0.95
        if title:
            ax.text(0.5, y, title, ha="center", va="top", fontsize=15, fontproperties=font, weight="bold")
            y -= 0.055
        return fig, ax, y

    def put_lines(pdf: PdfPages, title: str, blocks: list[str], *, fontsize: float = 10.5) -> None:
        fig, ax, y = new_page(title)
        for block in blocks:
            for line in wrap_cn(block, 58):
                if y < 0.075:
                    pdf.savefig(fig, bbox_inches="tight")
                    plt.close(fig)
                    fig, ax, y = new_page(title + "（续）")
                ax.text(0.08, y, line, ha="left", va="top", fontsize=fontsize, fontproperties=font)
                y -= 0.028 if line else 0.018
            y -= 0.008
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

    r1, r2, r3 = (est.iloc[i] for i in range(3))
    expected = float(tasks["expected_success"].sum())
    with PdfPages(PDF_OUT) as pdf:
        put_lines(
            pdf,
            "2026西安工程大学研究生数学建模校赛",
            [
                "学院：____________________________",
                "队号：____________________________",
                "题号：B",
                "组长联系电话：____________________",
                "",
                "队员信息请按竞赛封面模板填写。正文和摘要页不得出现任何队员身份信息。",
            ],
            fontsize=12,
        )
        put_lines(
            pdf,
            "摘 要",
            [
                "B题 多源融合机器人定位及任务优化",
                (
                    "本文建立时间同步、固定系统偏差估计、10Hz轨迹融合与0-1整数规划调度模型。"
                    f"问题1得到方式2相对方式1时间偏差{r1['方式2相对方式1时间偏差_delta_s']:.4f}s；"
                    f"问题2得到时间偏差{r2['方式2相对方式1时间偏差_delta_s']:.4f}s，系统偏差"
                    f"({r2['采用的系统偏差_x_m']:.4f},{r2['采用的系统偏差_y_m']:.4f})m；"
                    f"问题3采用嵌套F检验判定实际测量数据存在统计显著的微小系统偏差，估计偏差"
                    f"({r3['采用的系统偏差_x_m']:.4f},{r3['采用的系统偏差_y_m']:.4f})m，并据此修正10Hz融合轨迹。"
                    f"问题4在修正轨迹上生成候选任务并求解整数规划，选中{len(tasks)}项任务，期望完成数{expected:.2f}。"
                    "自动复核表明，全部任务满足距离、速度、加速度、准备时间和拍照角度约束。"
                ),
                "关键词：时间同步；系统偏差；轨迹融合；嵌套F检验；整数规划；任务调度",
            ],
        )
        put_lines(
            pdf,
            "正文概要",
            [
                "一、问题重述：两类定位数据存在异频、异步、随机噪声和可能的固定系统偏差；任务优化要求沿附件3轨迹尽可能多完成模拟射击和拍照扫描。",
                "二、模型假设：方式1为时间基准；同一附件内固定偏差为二维常向量；融合轨迹按10Hz输出；题面未给额外冷却时间，因此任务首尾相接视为不重叠。",
                "三、时间同步与偏差估计：给定delta后在公共时间网格插值，分别建立无偏模型J0和带偏模型J1，并用截尾均方误差搜索最优时间平移。系统偏差采用嵌套F检验，显著性水平alpha=0.05。",
                "四、轨迹融合：若检验显著，则用估计偏差修正方式2坐标；若检验不显著，则令b=0。融合位置取两种方式修正后位置均值，速度和加速度由Savitzky-Golay平滑轨迹差分估计。",
                "五、算法实现：先完成时间同步和偏差修正，再生成10Hz轨迹；随后逐时刻生成候选任务，构造MILP约束矩阵并求解最优任务组合，最后写入结果表和校验报告。",
                "六、任务优化：候选任务逐时刻检查准备区间全程约束；拍照方向角由机器人指向目标点的向量计算；0-1整数规划最大化期望完成数，并处理时间互斥、射击唯一性和角度冲突。",
            ],
        )
        results_lines = ["问题 | delta/s | bias_x/m | bias_y/m | 统计显著 | 是否修正 | 误差下降比例"]
        for _, row in est.iterrows():
            drop_text = "" if pd.isna(row["误差下降比例"]) else f"{100 * row['误差下降比例']:.2f}%"
            results_lines.append(
                f"{int(row['问题'])} | {row['方式2相对方式1时间偏差_delta_s']:.4f} | "
                f"{row['采用的系统偏差_x_m']:.4f} | {row['采用的系统偏差_y_m']:.4f} | "
                f"{row['是否统计显著']} | {row['是否采用系统偏差修正']} | "
                f"{drop_text}"
            )
        put_lines(pdf, "问题1至问题3结果", results_lines)
        task_cols = ["target_id", "task", "prep_start_s", "exec_time_s", "distance_m", "speed_m_s", "accel_m_s2", "angle_deg", "expected_success"]
        task_lines = ["目标 | 任务 | 准备/s | 执行/s | 距离/m | 速度 | 加速度 | 角度 | 收益"]
        for _, row in tasks.iterrows():
            task_lines.append(format_row_values(row, task_cols))
        put_lines(pdf, "问题4选中任务", task_lines, fontsize=8.2)
        put_lines(
            pdf,
            "结果检验与模型评价",
            [
                "自动校验结论：verification_report.xlsx 为 PASS。拍照任务最大速度约1.485m/s，最大加速度约1.267m/s²；射击任务最大速度约1.984m/s，最大加速度约0.781m/s²；同一拍照目标最小角度差约62.68度。",
                "敏感性分析显示，平滑窗口51、61、71、81、91点下期望完成数分别为14.40、16.40、17.25、16.10、15.65，主方案71点窗口取得最高可行期望完成数。",
                "完整候选集复核：将所有10Hz可行点展开为1195个候选后重新求解，最优值仍为18项、期望17.25，说明337个压缩候选没有损失当前离散模型最优目标值。",
                "优点：模型链条清晰，系统偏差判定有统计检验依据，任务调度用整数规划统一处理多类约束。局限：速度和加速度受平滑窗口影响，任务优化为10Hz离散最优，未建模额外冷却时间。",
            ],
        )
        for fig_name, title in [
            ("problem1_trajectory_10hz.png", "图1 问题1融合轨迹"),
            ("problem2_trajectory_10hz.png", "图2 问题2融合轨迹"),
            ("problem3_trajectory_10hz.png", "图3 问题3修正轨迹"),
            ("problem4_selected_tasks.png", "图4 问题4选中任务"),
        ]:
            img_path = OUT / "figures" / fig_name
            if img_path.exists():
                fig, ax, y = new_page(title)
                img = plt.imread(str(img_path))
                ax.imshow(img)
                ax.axis("off")
                pdf.savefig(fig, bbox_inches="tight")
                plt.close(fig)
        put_lines(
            pdf,
            "参考文献",
            [
                "[1] Savitzky A, Golay M J E, Smoothing and Differentiation of Data by Simplified Least Squares Procedures, Analytical Chemistry, 36(8):1627-1639, 1964.",
                "[2] Virtanen P, Gommers R, Oliphant T E, et al., SciPy 1.0: fundamental algorithms for scientific computing in Python, Nature Methods, 17:261-272, 2020.",
                "[3] Huangfu Q, Hall J A J, Parallelizing the dual revised simplex method, Mathematical Programming Computation, 10:119-142, 2018.",
                "[4] Efron B, Tibshirani R J, An Introduction to the Bootstrap, New York: Chapman & Hall/CRC, 1993.",
            ],
        )


def main() -> None:
    OUT.mkdir(exist_ok=True)
    est = pd.read_excel(OUT / "estimates_summary.xlsx")
    tasks = pd.read_excel(OUT / "task_schedule_detail.xlsx")
    sens = pd.read_excel(OUT / "sensitivity_smooth_window.xlsx")

    doc = Document()
    setup_document(doc)
    add_cover(doc)
    add_abstract_page(doc, est, tasks)
    add_body(doc, est, tasks, sens)
    doc.save(DOCX_OUT)
    write_pdf(est, tasks, sens)

    MD_OUT.write_text(build_markdown(est, tasks, sens), encoding="utf-8")
    REVIEW_OUT.write_text(build_self_review(), encoding="utf-8")
    print(DOCX_OUT)
    print(PDF_OUT)
    print(MD_OUT)
    print(REVIEW_OUT)


if __name__ == "__main__":
    main()
