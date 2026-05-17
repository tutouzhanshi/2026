# -*- coding: utf-8 -*-
"""问题4：基于问题3轨迹生成候选任务并求解任务调度。"""
from __future__ import annotations

from pathlib import Path
from typing import Iterable
import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix

from problem1 import DT, output_dir, project_root
from problem3 import solve_problem3


SHOOT_DISTANCE = (5.0, 30.0)
SHOOT_SPEED_MAX = 2.0
SHOOT_ACCEL_MAX = 1.5
SHOOT_PREP = 1.5
SHOOT_HIT_PROB = 0.85

PHOTO_DISTANCE = (10.0, 40.0)
PHOTO_SPEED_MAX = 1.5
PHOTO_ACCEL_MAX = 1.5
PHOTO_PREP = 0.5
PHOTO_ANGLE_MIN = 60.0


def rolling_all(mask: np.ndarray, window: int) -> np.ndarray:
    """判断每个采样点之前连续window个点是否全部满足约束。"""
    if window <= 1:
        return mask.astype(bool)
    out = np.zeros_like(mask, dtype=bool)
    csum = np.r_[0, np.cumsum(mask.astype(int))]
    out[window - 1 :] = (csum[window:] - csum[:-window]) == window
    return out


def angle_deg(vec: np.ndarray) -> np.ndarray:
    """计算二维向量的方向角，范围为[0, 360)。"""
    return (np.degrees(np.arctan2(vec[:, 1], vec[:, 0])) + 360.0) % 360.0


def angle_diff(a: float, b: float) -> float:
    """计算两个方向角的最小夹角。"""
    return float(abs((a - b + 180.0) % 360.0 - 180.0))


def continuous_runs(indices: np.ndarray) -> Iterable[np.ndarray]:
    """把连续可行采样点拆成若干连续区间。"""
    if len(indices) == 0:
        return
    start = 0
    for i in range(1, len(indices)):
        if indices[i] != indices[i - 1] + 1:
            yield indices[start:i]
            start = i
    yield indices[start:]


def select_representatives(indices: np.ndarray, score: np.ndarray, stride: int = 5) -> np.ndarray:
    """从连续可行区间中选少量代表点，控制MILP规模。"""
    keep: set[int] = set()
    for run in continuous_runs(indices):
        keep.add(int(run[0]))
        keep.add(int(run[-1]))
        keep.add(int(run[np.argmin(score[run])]))
        for idx in run[::stride]:
            keep.add(int(idx))
    return np.array(sorted(keep), dtype=int)


def build_task_candidates(traj: pd.DataFrame, target_path: Path) -> pd.DataFrame:
    """根据轨迹和题目约束生成所有单项可行的候选任务。"""
    t = traj["time_s"].to_numpy(float)
    xy = traj[["x_m", "y_m"]].to_numpy(float)
    speed = traj["speed_m_s"].to_numpy(float)
    accel = traj["accel_m_s2"].to_numpy(float)
    shots = pd.read_excel(target_path, sheet_name=0).dropna().iloc[:, 0:3]
    photos = pd.read_excel(target_path, sheet_name=1).dropna().iloc[:, 0:3]
    rows: list[dict[str, object]] = []

    shoot_window = int(round(SHOOT_PREP / DT)) + 1
    for _, row in shots.iterrows():
        target_id = str(row.iloc[0])
        point = row.iloc[1:3].to_numpy(float)
        dist = np.linalg.norm(xy - point, axis=1)
        feasible = (
            (dist >= SHOOT_DISTANCE[0])
            & (dist <= SHOOT_DISTANCE[1])
            & (speed <= SHOOT_SPEED_MAX)
            & (accel <= SHOOT_ACCEL_MAX)
        )
        for idx in select_representatives(np.where(rolling_all(feasible, shoot_window))[0], dist):
            rows.append(
                {
                    "target_id": target_id,
                    "task": "模拟射击",
                    "prep_start_s": round(float(t[idx] - SHOOT_PREP), 1),
                    "exec_time_s": round(float(t[idx]), 1),
                    "distance_m": float(dist[idx]),
                    "speed_m_s": float(speed[idx]),
                    "accel_m_s2": float(accel[idx]),
                    "angle_deg": np.nan,
                    "expected_success": SHOOT_HIT_PROB,
                }
            )

    photo_window = int(round(PHOTO_PREP / DT)) + 1
    for _, row in photos.iterrows():
        target_id = str(row.iloc[0])
        point = row.iloc[1:3].to_numpy(float)
        vec = point - xy
        dist = np.linalg.norm(vec, axis=1)
        angle = angle_deg(vec)
        feasible = (
            (dist >= PHOTO_DISTANCE[0])
            & (dist <= PHOTO_DISTANCE[1])
            & (speed <= PHOTO_SPEED_MAX)
            & (accel <= PHOTO_ACCEL_MAX)
        )
        idxs = np.where(rolling_all(feasible, photo_window))[0]
        reps = set(select_representatives(idxs, dist).tolist())
        for bin_id in range(24):
            inside = idxs[(angle[idxs] >= bin_id * 15.0) & (angle[idxs] < (bin_id + 1) * 15.0)]
            if len(inside):
                reps.add(int(inside[np.argmin(dist[inside])]))
        for idx in sorted(reps):
            rows.append(
                {
                    "target_id": target_id,
                    "task": "拍照",
                    "prep_start_s": round(float(t[idx] - PHOTO_PREP), 1),
                    "exec_time_s": round(float(t[idx]), 1),
                    "distance_m": float(dist[idx]),
                    "speed_m_s": float(speed[idx]),
                    "accel_m_s2": float(accel[idx]),
                    "angle_deg": float(angle[idx]),
                    "expected_success": 1.0,
                }
            )

    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df = df.drop_duplicates(["target_id", "task", "prep_start_s", "exec_time_s"])
    return df.sort_values(["exec_time_s", "prep_start_s", "target_id", "task"]).reset_index(drop=True)


def schedule_tasks_milp(candidates: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    """在候选任务集合上求解0-1整数规划。"""
    if candidates.empty:
        return candidates, "empty"
    df = candidates.copy().reset_index(drop=True)
    n = len(df)
    starts = df["prep_start_s"].to_numpy(float)
    ends = df["exec_time_s"].to_numpy(float)
    weights = df["expected_success"].to_numpy(float)

    rows: list[list[int]] = []
    for _target_id, group in df[df["task"] == "模拟射击"].groupby("target_id"):
        if len(group.index) > 1:
            rows.append(group.index.to_list())

    order = np.argsort(starts)
    for pos, i in enumerate(order):
        for j in order[pos + 1 :]:
            if starts[j] >= ends[i] - 1e-9:
                break
            if starts[i] < ends[j] - 1e-9 and starts[j] < ends[i] - 1e-9:
                rows.append([int(i), int(j)])

    photos = df[df["task"] == "拍照"]
    for _target_id, group in photos.groupby("target_id"):
        idxs = group.index.to_list()
        angles = df.loc[idxs, "angle_deg"].to_numpy(float)
        for a in range(len(idxs)):
            for b in range(a + 1, len(idxs)):
                if angle_diff(float(angles[a]), float(angles[b])) < PHOTO_ANGLE_MIN - 1e-9:
                    rows.append([idxs[a], idxs[b]])

    if rows:
        matrix = lil_matrix((len(rows), n), dtype=float)
        for r, cols in enumerate(rows):
            for c in cols:
                matrix[r, c] = 1.0
        constraints = LinearConstraint(matrix.tocsr(), np.full(len(rows), -np.inf), np.ones(len(rows)))
    else:
        constraints = ()

    # 主目标是最大化期望完成数；后两项只用于稳定多个同等最优解的输出顺序。
    cost = -weights + 1e-7 * starts + 1e-8 * df["distance_m"].to_numpy(float)
    res = milp(
        c=cost,
        integrality=np.ones(n),
        bounds=Bounds(np.zeros(n), np.ones(n)),
        constraints=constraints,
        options={"time_limit": 90, "mip_rel_gap": 0.0},
    )
    if not res.success:
        raise RuntimeError(f"MILP failed: {res.message}")
    chosen = np.where(res.x > 0.5)[0]
    out = df.iloc[chosen].copy().sort_values(["prep_start_s", "exec_time_s"]).reset_index(drop=True)
    return out, "milp_optimal"


def write_result_xlsx(template: Path, output: Path, tasks: pd.DataFrame) -> None:
    """把选中任务写入题目给定的result.xlsx格式。"""
    if template.resolve() != output.resolve():
        shutil.copy2(template, output)
    wb = load_workbook(output)
    ws = wb.active
    for r in range(2, max(ws.max_row + 1, len(tasks) + 3)):
        for c in range(1, 6):
            ws.cell(r, c).value = None
    for i, row in tasks.iterrows():
        r = i + 2
        values = [i + 1, row["target_id"], row["task"], float(row["prep_start_s"]), float(row["exec_time_s"])]
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="宋体", size=12)
    wb.save(output)


def verify_schedule(traj: pd.DataFrame, tasks: pd.DataFrame, target_path: Path) -> list[str]:
    """复核选中任务是否满足运动学、时间互斥和目标约束。"""
    issues: list[str] = []
    shots = pd.read_excel(target_path, sheet_name=0).dropna().iloc[:, 0:3]
    photos = pd.read_excel(target_path, sheet_name=1).dropna().iloc[:, 0:3]
    points = {
        str(row.iloc[0]): row.iloc[1:3].to_numpy(float)
        for _, row in pd.concat([shots, photos], ignore_index=True).iterrows()
    }

    last_end = -float("inf")
    for _, row in tasks.sort_values("prep_start_s").iterrows():
        start = float(row["prep_start_s"])
        end = float(row["exec_time_s"])
        if start < last_end - 1e-9:
            issues.append(f"time overlap near {row['target_id']}")
        last_end = max(last_end, end)

        seg = traj[(traj["time_s"] >= start - 1e-9) & (traj["time_s"] <= end + 1e-9)]
        xy = seg[["x_m", "y_m"]].to_numpy(float)
        dist = np.linalg.norm(xy - points[str(row["target_id"])], axis=1)
        speed = seg["speed_m_s"].to_numpy(float)
        accel = seg["accel_m_s2"].to_numpy(float)
        if row["task"] == "模拟射击":
            ok = (
                (dist >= SHOOT_DISTANCE[0])
                & (dist <= SHOOT_DISTANCE[1])
                & (speed <= SHOOT_SPEED_MAX)
                & (accel <= SHOOT_ACCEL_MAX)
            )
        else:
            ok = (
                (dist >= PHOTO_DISTANCE[0])
                & (dist <= PHOTO_DISTANCE[1])
                & (speed <= PHOTO_SPEED_MAX)
                & (accel <= PHOTO_ACCEL_MAX)
            )
        if len(ok) == 0 or not bool(np.all(ok)):
            issues.append(f"constraint violation: {row['target_id']} {row['task']}")

    for target_id, group in tasks[tasks["task"] == "拍照"].groupby("target_id"):
        angles = group["angle_deg"].dropna().to_numpy(float)
        for i in range(len(angles)):
            for j in range(i + 1, len(angles)):
                if angle_diff(float(angles[i]), float(angles[j])) < PHOTO_ANGLE_MIN - 1e-9:
                    issues.append(f"photo angle conflict: {target_id}")
    if tasks[tasks["task"] == "模拟射击"]["target_id"].duplicated().any():
        issues.append("duplicated shooting target")
    return issues


def solve_problem4(root: Path | None = None, write_outputs: bool = True) -> tuple[pd.DataFrame, str, list[str]]:
    """求解问题4并写出任务候选、任务明细、result.xlsx和复核结果。"""
    root = root or project_root()
    _alignment, trajectory = solve_problem3(root, write_outputs=False)
    candidates = build_task_candidates(trajectory, root / "附件4.xlsx")
    tasks, status = schedule_tasks_milp(candidates)
    issues = verify_schedule(trajectory, tasks, root / "附件4.xlsx")

    if write_outputs:
        out = output_dir(root)
        trajectory.to_excel(out / "problem3_10Hz_trajectory.xlsx", index=False)
        candidates.to_excel(out / "task_candidates.xlsx", index=False)
        tasks.to_excel(out / "task_schedule_detail.xlsx", index=False)
        pd.DataFrame({"status": [status], "check": ["PASS" if not issues else "FAIL"]}).to_excel(
            out / "verification_report.xlsx", index=False
        )
        write_result_xlsx(root / "result.xlsx", out / "result.xlsx", tasks)
        write_result_xlsx(root / "result.xlsx", root / "result.xlsx", tasks)
    return tasks, status, issues


def main() -> None:
    tasks, status, issues = solve_problem4()
    print(f"problem4 status = {status}")
    print(f"selected tasks = {len(tasks)}, expected = {tasks['expected_success'].sum():.4f}")
    print("check = PASS" if not issues else "check = FAIL")
    for issue in issues:
        print(issue)


if __name__ == "__main__":
    main()
